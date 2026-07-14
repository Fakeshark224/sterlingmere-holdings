import smtplib
from email.mime.multipart import MIMEMultipart
from email.mime.text import MIMEText
import openpyxl
import os
from datetime import datetime
import time
import csv

# --- CONFIGURATION ---
SMTP_SERVER = "smtp.gmail.com"
SMTP_PORT = 587
EMAIL_ADDRESS = "sterlingmeremail@gmail.com"
EMAIL_PASSWORD = "lnxmzacxtkfdrbxb"  # Note: You may need to use an App Password if 2FA is enabled.
CAMPAIGN_FILE = "campaign_tracking.csv"
LEADS_FILE = "leads.csv"
FOLLOWUP_DAYS = 3

def get_excel_filename():
    now = datetime.now()
    month_name = now.strftime("%B")
    year = now.strftime("%Y")
    return f"Outreach_{month_name}_{year}.xlsx"

def get_sheet_name():
    now = datetime.now()
    return now.strftime("%d-%m-%Y")

def setup_excel_logger():
    filename = get_excel_filename()
    sheet_name = get_sheet_name()
    if os.path.exists(filename):
        wb = openpyxl.load_workbook(filename)
    else:
        wb = openpyxl.Workbook()
        if 'Sheet' in wb.sheetnames:
            del wb['Sheet']
    if sheet_name not in wb.sheetnames:
        ws = wb.create_sheet(title=sheet_name)
        ws.append(["Timestamp", "Business Name", "Email", "Niche", "Status", "Notes"])
        for cell in ws[1]:
            cell.font = openpyxl.styles.Font(bold=True)
    else:
        ws = wb[sheet_name]
    return wb, ws, filename

def log_to_excel(wb, ws, filename, business_name, email, niche, status, notes=""):
    timestamp = datetime.now().strftime("%Y-%m-%d %H:%M:%S")
    ws.append([timestamp, business_name, email, niche, status, notes])
    wb.save(filename)

def send_email(to_email, business_name, niche, is_followup=False):
    """Sends a personalized plain text email to the target."""
    try:
        template_file = "followup_template.txt" if is_followup else "email_template.txt"
        with open(template_file, "r", encoding="utf-8") as f:
            text_content = f.read()
            
        text_content = text_content.replace("{business_name}", business_name)
        text_content = text_content.replace("{niche}", niche)
        
        msg = MIMEMultipart()
        msg['From'] = f"Sterlingmere Holdings <{EMAIL_ADDRESS}>"
        msg['To'] = to_email
        msg['Subject'] = f"Quick question about {business_name}" if is_followup else f"Upgrading {business_name}'s Digital Presence"
        
        msg.attach(MIMEText(text_content, 'plain'))
        
        server = smtplib.SMTP(SMTP_SERVER, SMTP_PORT)
        server.starttls()
        server.login(EMAIL_ADDRESS, EMAIL_PASSWORD)
        server.send_message(msg)
        server.quit()
        return True, "Success"
    except Exception as e:
        return False, str(e)

def load_campaigns():
    campaigns = {}
    if os.path.exists(CAMPAIGN_FILE):
        with open(CAMPAIGN_FILE, 'r', encoding='utf-8') as f:
            reader = csv.DictReader(f)
            for row in reader:
                email = row.get('Email', '').lower().strip()
                if email:
                    campaigns[email] = row
    return campaigns

def save_campaigns(campaigns):
    with open(CAMPAIGN_FILE, 'w', newline='', encoding='utf-8') as f:
        writer = csv.DictWriter(f, fieldnames=["Email", "Business Name", "Niche", "Stage", "Last Emailed Date"])
        writer.writeheader()
        for email, data in campaigns.items():
            writer.writerow(data)

def main():
    print(f"Starting Sterlingmere Outreach Bot (Follow-Up Engine Active)...")
    wb, ws, filename = setup_excel_logger()
    print(f"Logging data to: {filename}, Sheet: {ws.title}")
    
    # Port over old sent_emails.txt if it exists
    campaigns = load_campaigns()
    if os.path.exists("sent_emails.txt"):
        with open("sent_emails.txt", "r", encoding="utf-8") as f:
            for line in f:
                email = line.strip().lower()
                if email and email not in campaigns:
                    # Mark as Completed so they are skipped in the future but don't get follow-ups without names
                    campaigns[email] = {
                        "Email": email, "Business Name": "Unknown", "Niche": "Unknown", 
                        "Stage": "Completed", "Last Emailed Date": today.strftime("%Y-%m-%d")
                    }
        # Rename so we don't do this again
        os.rename("sent_emails.txt", "sent_emails_old.txt")
    
    today = datetime.now()
    today_str = today.strftime("%Y-%m-%d")
    
    # 1. PROCESS FOLLOW UPS
    print("\n[1] Checking for Follow-Ups...")
    followups_sent = 0
    for email, data in campaigns.items():
        if data['Stage'] == '1':
            last_date_str = data.get('Last Emailed Date', '2020-01-01')
            try:
                last_date = datetime.strptime(last_date_str, "%Y-%m-%d")
                days_since = (today - last_date).days
            except:
                days_since = 999
                
            if days_since >= FOLLOWUP_DAYS:
                b_name = data.get('Business Name', 'there')
                niche = data.get('Niche', 'your industry')
                print(f"Sending FOLLOW-UP to {b_name} ({email})...")
                
                success, message = send_email(email, b_name, niche, is_followup=True)
                status = "Follow-Up Sent" if success else "Follow-Up Failed"
                print(f"Result: {status} - {message}")
                log_to_excel(wb, ws, filename, b_name, email, niche, status, message)
                
                if success:
                    data['Stage'] = '2'
                    data['Last Emailed Date'] = today_str
                    followups_sent += 1
                time.sleep(2)
                
    if followups_sent == 0:
        print("No follow-ups due today.")
        
    # 2. PROCESS NEW LEADS
    print("\n[2] Checking for New Leads...")
    if not os.path.exists(LEADS_FILE):
        print(f"No {LEADS_FILE} found. Skipping new leads.")
    else:
        new_leads_sent = 0
        with open(LEADS_FILE, mode='r', encoding='utf-8') as file:
            reader = csv.DictReader(file)
            for row in reader:
                b_name = row.get('Business Name', '')
                email = row.get('Email', '')
                niche = row.get('Niche', '')
                
                if not email or not b_name:
                    continue
                    
                clean_email = email.lower().strip()
                if clean_email in campaigns:
                    print(f"Skipping {b_name} ({clean_email}) - Already in campaign tracker.")
                    log_to_excel(wb, ws, filename, b_name, email, niche, "Skipped", "Already in campaign tracker")
                    continue
                    
                print(f"Sending INITIAL PITCH to {b_name} ({clean_email})...")
                success, message = send_email(clean_email, b_name, niche, is_followup=False)
                status = "Initial Pitch Sent" if success else "Failed"
                print(f"Result: {status} - {message}")
                log_to_excel(wb, ws, filename, b_name, email, niche, status, message)
                
                if success:
                    campaigns[clean_email] = {
                        "Email": clean_email, "Business Name": b_name, "Niche": niche, 
                        "Stage": "1", "Last Emailed Date": today_str
                    }
                    new_leads_sent += 1
                time.sleep(2)
                
        # Clear the leads.csv file
        with open(LEADS_FILE, mode='w', newline='', encoding='utf-8') as file:
            writer = csv.writer(file)
            writer.writerow(["Business Name", "Email", "Niche", "Website"])
            
    # Save the campaign tracker
    save_campaigns(campaigns)
    print(f"\nOutreach complete! Data saved to {filename} and leads.csv queue cleared.")

if __name__ == "__main__":
    main()
